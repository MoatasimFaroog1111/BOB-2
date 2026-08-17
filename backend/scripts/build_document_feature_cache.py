#!/usr/bin/env python3
"""Build a resumable, SHA-addressed feature cache for Odoo training documents.

Design goals:
- read-only against the exported dataset;
- no permanent page-image copies;
- embedded PDF text first, OCR only when a page lacks useful text;
- adaptive Arabic+English Tesseract OCR for scans/images;
- one extraction per unique SHA-256 document;
- compressed SQLite persistence for Train Once -> Persist -> Load & Predict;
- resumable and deterministic.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import logging
import math
import re
import shutil
import sqlite3
import statistics
import sys
import zipfile
import zlib
from abc import ABC, abstractmethod
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree as ET

from PIL import Image, ImageOps
from pypdf import PdfReader
import pypdfium2 as pdfium
import pytesseract
from pytesseract import Output

logging.getLogger("pypdf").setLevel(logging.ERROR)

FEATURE_VERSION = "document-features-v1.0.0"
EMBEDDED_PAGE_MIN_CHARS = 40
MAX_STRUCTURED_TEXT_CHARS = 500_000


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def normalize_text(text: str | None) -> str:
    text = (text or "").replace("\x00", " ")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def compact_char_count(text: str) -> int:
    return len("".join(text.split()))


@dataclass(frozen=True)
class DocumentRef:
    sha256: str
    path: Path
    relative_path: str
    mime: str
    name: str


@dataclass
class ExtractionResult:
    status: str
    extractor: str
    text: str = ""
    page_count: int = 0
    embedded_pages: int = 0
    ocr_pages: int = 0
    ocr_mean_conf: float | None = None
    ocr_psm_counts: dict[int, int] = field(default_factory=dict)
    error: str | None = None


@dataclass
class OcrResult:
    text: str
    confidence: float
    psm: int


class AdaptiveOcrEngine:
    """Arabic+English OCR with a fast first pass and confidence-based fallback."""

    def __init__(self, timeout_seconds: int = 90) -> None:
        self.timeout_seconds = timeout_seconds

    @staticmethod
    def _prepare(image: Image.Image) -> Image.Image:
        image = ImageOps.exif_transpose(image).convert("L")
        image = ImageOps.autocontrast(image)

        width, height = image.size
        longest = max(width, height)

        # Upscale small captures for OCR; cap very large captures to control memory.
        if longest < 1800 and longest > 0:
            scale = 1800 / longest
            image = image.resize(
                (max(1, round(width * scale)), max(1, round(height * scale))),
                Image.Resampling.LANCZOS,
            )
        elif longest > 3400:
            scale = 3400 / longest
            image = image.resize(
                (max(1, round(width * scale)), max(1, round(height * scale))),
                Image.Resampling.LANCZOS,
            )

        return image

    def _run_psm(self, image: Image.Image, psm: int) -> OcrResult:
        data = pytesseract.image_to_data(
            image,
            lang="ara+eng",
            config=f"--oem 3 --psm {psm}",
            output_type=Output.DICT,
            timeout=self.timeout_seconds,
        )

        grouped: dict[tuple[int, int, int], list[str]] = defaultdict(list)
        weighted_conf_sum = 0.0
        weighted_char_sum = 0

        texts = data.get("text", [])
        confs = data.get("conf", [])
        blocks = data.get("block_num", [])
        pars = data.get("par_num", [])
        lines = data.get("line_num", [])

        for i, raw in enumerate(texts):
            token = normalize_text(str(raw))
            if not token:
                continue

            key = (
                int(blocks[i]) if i < len(blocks) else 0,
                int(pars[i]) if i < len(pars) else 0,
                int(lines[i]) if i < len(lines) else i,
            )
            grouped[key].append(token)

            try:
                conf = float(confs[i]) if i < len(confs) else -1.0
            except (TypeError, ValueError):
                conf = -1.0

            if conf >= 0:
                weight = max(1, len(token))
                weighted_conf_sum += conf * weight
                weighted_char_sum += weight

        lines_out = [" ".join(grouped[k]) for k in sorted(grouped)]
        text = normalize_text("\n".join(lines_out))
        confidence = (
            weighted_conf_sum / weighted_char_sum if weighted_char_sum else 0.0
        )
        return OcrResult(text=text, confidence=confidence, psm=psm)

    @staticmethod
    def _score(result: OcrResult) -> float:
        chars = compact_char_count(result.text)
        # Confidence dominates; useful text volume breaks close ties.
        return result.confidence + min(chars, 1500) / 150.0

    def extract(self, image: Image.Image) -> OcrResult:
        prepared = self._prepare(image)
        try:
            candidates: list[OcrResult] = []
            first_error: Exception | None = None

            try:
                first = self._run_psm(prepared, 3)
                candidates.append(first)
            except Exception as exc:  # Tesseract timeout/read error
                first_error = exc
                first = None

            # Fast path for clean documents.
            if (
                first is not None
                and first.confidence >= 70.0
                and compact_char_count(first.text) >= 120
            ):
                return first

            for psm in (6, 11):
                try:
                    candidates.append(self._run_psm(prepared, psm))
                except Exception as exc:
                    if first_error is None:
                        first_error = exc

            if not candidates:
                raise RuntimeError(f"All OCR modes failed: {first_error}")

            return max(candidates, key=self._score)
        finally:
            prepared.close()


class DocumentExtractor(ABC):
    @abstractmethod
    def extract(self, document: DocumentRef) -> ExtractionResult:
        raise NotImplementedError


class PdfDocumentExtractor(DocumentExtractor):
    def __init__(self, ocr: AdaptiveOcrEngine, render_scale: float = 2.5) -> None:
        self.ocr = ocr
        self.render_scale = render_scale

    def _ocr_page(self, pdf: Any, page_number: int) -> OcrResult:
        page = pdf[page_number]
        bitmap = None
        image = None
        try:
            bitmap = page.render(scale=self.render_scale, rotation=0)
            image = bitmap.to_pil()
            return self.ocr.extract(image)
        finally:
            if image is not None:
                image.close()
            if bitmap is not None and hasattr(bitmap, "close"):
                bitmap.close()
            if hasattr(page, "close"):
                page.close()

    def _ocr_all(self, path: Path) -> ExtractionResult:
        pdf = pdfium.PdfDocument(str(path))
        texts: list[str] = []
        confs: list[float] = []
        psm_counts: Counter[int] = Counter()
        try:
            for pno in range(len(pdf)):
                ocr = self._ocr_page(pdf, pno)
                texts.append(ocr.text)
                confs.append(ocr.confidence)
                psm_counts[ocr.psm] += 1
        finally:
            pdf.close()

        return ExtractionResult(
            status="success",
            extractor="pdf_ocr",
            text=normalize_text("\n\n".join(texts)),
            page_count=len(texts),
            ocr_pages=len(texts),
            ocr_mean_conf=statistics.fmean(confs) if confs else None,
            ocr_psm_counts=dict(psm_counts),
        )

    def extract(self, document: DocumentRef) -> ExtractionResult:
        try:
            reader = PdfReader(str(document.path), strict=False)
            page_count = len(reader.pages)
        except Exception:
            return self._ocr_all(document.path)

        texts: list[str] = []
        embedded_pages = 0
        ocr_pages = 0
        confs: list[float] = []
        psm_counts: Counter[int] = Counter()
        pdfium_doc = None

        try:
            for pno in range(page_count):
                embedded = ""
                try:
                    embedded = normalize_text(reader.pages[pno].extract_text() or "")
                except Exception:
                    embedded = ""

                if compact_char_count(embedded) >= EMBEDDED_PAGE_MIN_CHARS:
                    texts.append(embedded)
                    embedded_pages += 1
                    continue

                if pdfium_doc is None:
                    pdfium_doc = pdfium.PdfDocument(str(document.path))

                ocr = self._ocr_page(pdfium_doc, pno)
                texts.append(ocr.text)
                ocr_pages += 1
                confs.append(ocr.confidence)
                psm_counts[ocr.psm] += 1
        finally:
            if pdfium_doc is not None:
                pdfium_doc.close()

        if ocr_pages == 0:
            mode = "pdf_text"
        elif embedded_pages == 0:
            mode = "pdf_ocr"
        else:
            mode = "pdf_hybrid"

        return ExtractionResult(
            status="success",
            extractor=mode,
            text=normalize_text("\n\n".join(texts)),
            page_count=page_count,
            embedded_pages=embedded_pages,
            ocr_pages=ocr_pages,
            ocr_mean_conf=statistics.fmean(confs) if confs else None,
            ocr_psm_counts=dict(psm_counts),
        )


class ImageDocumentExtractor(DocumentExtractor):
    def __init__(self, ocr: AdaptiveOcrEngine) -> None:
        self.ocr = ocr

    def extract(self, document: DocumentRef) -> ExtractionResult:
        with Image.open(document.path) as image:
            ocr = self.ocr.extract(image)
        return ExtractionResult(
            status="success",
            extractor="image_ocr",
            text=ocr.text,
            page_count=1,
            ocr_pages=1,
            ocr_mean_conf=ocr.confidence,
            ocr_psm_counts={ocr.psm: 1},
        )


class PlainTextDocumentExtractor(DocumentExtractor):
    def extract(self, document: DocumentRef) -> ExtractionResult:
        raw = document.path.read_bytes()
        text = ""
        for encoding in ("utf-8-sig", "utf-16", "cp1256", "latin-1"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        return ExtractionResult(
            status="success",
            extractor="plain_text",
            text=normalize_text(text)[:MAX_STRUCTURED_TEXT_CHARS],
            page_count=1,
        )


class XlsxDocumentExtractor(DocumentExtractor):
    def extract(self, document: DocumentRef) -> ExtractionResult:
        try:
            import openpyxl
        except ImportError as exc:
            raise RuntimeError("openpyxl is required for .xlsx extraction") from exc

        workbook = openpyxl.load_workbook(
            document.path,
            read_only=True,
            data_only=True,
        )
        parts: list[str] = []
        chars = 0
        try:
            for sheet in workbook.worksheets:
                parts.append(f"[SHEET] {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    values = [normalize_text(str(v)) for v in row if v is not None]
                    if values:
                        line = " | ".join(values)
                        parts.append(line)
                        chars += len(line)
                    if chars >= MAX_STRUCTURED_TEXT_CHARS:
                        break
                if chars >= MAX_STRUCTURED_TEXT_CHARS:
                    break
        finally:
            workbook.close()

        return ExtractionResult(
            status="success",
            extractor="xlsx_cells",
            text=normalize_text("\n".join(parts))[:MAX_STRUCTURED_TEXT_CHARS],
            page_count=max(1, len(parts)),
        )


class XlsDocumentExtractor(DocumentExtractor):
    def extract(self, document: DocumentRef) -> ExtractionResult:
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError("xlrd is required for legacy .xls extraction") from exc

        book = xlrd.open_workbook(str(document.path), on_demand=True)
        parts: list[str] = []
        chars = 0
        try:
            for sheet_name in book.sheet_names():
                sheet = book.sheet_by_name(sheet_name)
                parts.append(f"[SHEET] {sheet_name}")
                for r in range(sheet.nrows):
                    values = [
                        normalize_text(str(sheet.cell_value(r, c)))
                        for c in range(sheet.ncols)
                        if sheet.cell_value(r, c) not in (None, "")
                    ]
                    if values:
                        line = " | ".join(values)
                        parts.append(line)
                        chars += len(line)
                    if chars >= MAX_STRUCTURED_TEXT_CHARS:
                        break
                if chars >= MAX_STRUCTURED_TEXT_CHARS:
                    break
        finally:
            book.release_resources()

        return ExtractionResult(
            status="success",
            extractor="xls_cells",
            text=normalize_text("\n".join(parts))[:MAX_STRUCTURED_TEXT_CHARS],
            page_count=max(1, len(book.sheet_names())),
        )


class DocxDocumentExtractor(DocumentExtractor):
    def extract(self, document: DocumentRef) -> ExtractionResult:
        with zipfile.ZipFile(document.path) as archive:
            xml = archive.read("word/document.xml")
        root = ET.fromstring(xml)
        texts = [node.text or "" for node in root.iter() if node.tag.endswith("}t")]
        return ExtractionResult(
            status="success",
            extractor="docx_xml",
            text=normalize_text(" ".join(texts))[:MAX_STRUCTURED_TEXT_CHARS],
            page_count=1,
        )


class ZipMetadataExtractor(DocumentExtractor):
    def extract(self, document: DocumentRef) -> ExtractionResult:
        with zipfile.ZipFile(document.path) as archive:
            rows = [
                f"{item.filename} | {item.file_size} bytes"
                for item in archive.infolist()[:5000]
            ]
        return ExtractionResult(
            status="success",
            extractor="zip_metadata",
            text=normalize_text("\n".join(rows))[:MAX_STRUCTURED_TEXT_CHARS],
            page_count=1,
        )


class UnsupportedDocumentExtractor(DocumentExtractor):
    def extract(self, document: DocumentRef) -> ExtractionResult:
        return ExtractionResult(
            status="unsupported",
            extractor="unsupported",
            text="",
            page_count=0,
            error=f"Unsupported extension: {document.path.suffix.lower()}",
        )


class ExtractorRegistry:
    def __init__(self, ocr: AdaptiveOcrEngine) -> None:
        self.pdf = PdfDocumentExtractor(ocr)
        self.image = ImageDocumentExtractor(ocr)
        self.plain = PlainTextDocumentExtractor()
        self.xlsx = XlsxDocumentExtractor()
        self.xls = XlsDocumentExtractor()
        self.docx = DocxDocumentExtractor()
        self.zip = ZipMetadataExtractor()
        self.unsupported = UnsupportedDocumentExtractor()

    def resolve(self, path: Path) -> DocumentExtractor:
        ext = path.suffix.lower()
        if ext == ".pdf":
            return self.pdf
        if ext in {".jpg", ".jpeg", ".png"}:
            return self.image
        if ext in {".txt", ".csv", ".xml"}:
            return self.plain
        if ext == ".xlsx":
            return self.xlsx
        if ext == ".xls":
            return self.xls
        if ext == ".docx":
            return self.docx
        if ext == ".zip":
            return self.zip
        return self.unsupported


class FeatureCacheRepository:
    def __init__(self, db_path: Path) -> None:
        db_path.parent.mkdir(parents=True, exist_ok=True)
        self.db_path = db_path
        self.db = sqlite3.connect(db_path)
        self.db.execute("PRAGMA synchronous=NORMAL")
        self.db.execute(
            """
            CREATE TABLE IF NOT EXISTS document_features (
                sha256 TEXT PRIMARY KEY,
                feature_version TEXT NOT NULL,
                relative_path TEXT NOT NULL,
                mime TEXT,
                name TEXT,
                extension TEXT,
                size_bytes INTEGER NOT NULL,
                status TEXT NOT NULL,
                extractor TEXT NOT NULL,
                page_count INTEGER NOT NULL DEFAULT 0,
                embedded_pages INTEGER NOT NULL DEFAULT 0,
                ocr_pages INTEGER NOT NULL DEFAULT 0,
                ocr_mean_conf REAL,
                ocr_psm_counts_json TEXT NOT NULL DEFAULT '{}',
                char_count INTEGER NOT NULL DEFAULT 0,
                line_count INTEGER NOT NULL DEFAULT 0,
                arabic_char_count INTEGER NOT NULL DEFAULT 0,
                latin_char_count INTEGER NOT NULL DEFAULT 0,
                digit_char_count INTEGER NOT NULL DEFAULT 0,
                text_sha256 TEXT,
                text_zlib BLOB,
                error TEXT,
                updated_at TEXT NOT NULL
            )
            """
        )
        self.db.execute(
            "CREATE INDEX IF NOT EXISTS ix_document_features_status ON document_features(status)"
        )
        self.db.execute(
            "CREATE INDEX IF NOT EXISTS ix_document_features_extractor ON document_features(extractor)"
        )
        self.db.commit()

    def is_current_success(self, sha256: str) -> bool:
        row = self.db.execute(
            "SELECT status, feature_version FROM document_features WHERE sha256=?",
            (sha256,),
        ).fetchone()
        return bool(row and row[0] == "success" and row[1] == FEATURE_VERSION)

    def save(self, document: DocumentRef, result: ExtractionResult) -> None:
        text = normalize_text(result.text)
        arabic_chars = len(re.findall(r"[\u0600-\u06FF]", text))
        latin_chars = len(re.findall(r"[A-Za-z]", text))
        digit_chars = len(re.findall(r"[0-9\u0660-\u0669]", text))
        line_count = len([line for line in text.splitlines() if line.strip()])
        raw = text.encode("utf-8")
        compressed = zlib.compress(raw, level=6) if raw else None
        text_sha = hashlib.sha256(raw).hexdigest() if raw else None
        size_bytes = document.path.stat().st_size if document.path.exists() else 0

        self.db.execute(
            """
            INSERT INTO document_features (
                sha256, feature_version, relative_path, mime, name, extension,
                size_bytes, status, extractor, page_count, embedded_pages,
                ocr_pages, ocr_mean_conf, ocr_psm_counts_json, char_count,
                line_count, arabic_char_count, latin_char_count,
                digit_char_count, text_sha256, text_zlib, error, updated_at
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            ON CONFLICT(sha256) DO UPDATE SET
                feature_version=excluded.feature_version,
                relative_path=excluded.relative_path,
                mime=excluded.mime,
                name=excluded.name,
                extension=excluded.extension,
                size_bytes=excluded.size_bytes,
                status=excluded.status,
                extractor=excluded.extractor,
                page_count=excluded.page_count,
                embedded_pages=excluded.embedded_pages,
                ocr_pages=excluded.ocr_pages,
                ocr_mean_conf=excluded.ocr_mean_conf,
                ocr_psm_counts_json=excluded.ocr_psm_counts_json,
                char_count=excluded.char_count,
                line_count=excluded.line_count,
                arabic_char_count=excluded.arabic_char_count,
                latin_char_count=excluded.latin_char_count,
                digit_char_count=excluded.digit_char_count,
                text_sha256=excluded.text_sha256,
                text_zlib=excluded.text_zlib,
                error=excluded.error,
                updated_at=excluded.updated_at
            """,
            (
                document.sha256,
                FEATURE_VERSION,
                document.relative_path,
                document.mime,
                document.name,
                document.path.suffix.lower(),
                size_bytes,
                result.status,
                result.extractor,
                result.page_count,
                result.embedded_pages,
                result.ocr_pages,
                result.ocr_mean_conf,
                json.dumps(result.ocr_psm_counts, sort_keys=True),
                len(text),
                line_count,
                arabic_chars,
                latin_chars,
                digit_chars,
                text_sha,
                compressed,
                result.error,
                utc_now(),
            ),
        )

    def commit(self) -> None:
        self.db.commit()

    def summary(self) -> dict[str, Any]:
        status_counts = dict(
            self.db.execute(
                "SELECT status, COUNT(*) FROM document_features WHERE feature_version=? GROUP BY status",
                (FEATURE_VERSION,),
            ).fetchall()
        )
        extractor_counts = dict(
            self.db.execute(
                "SELECT extractor, COUNT(*) FROM document_features WHERE feature_version=? GROUP BY extractor",
                (FEATURE_VERSION,),
            ).fetchall()
        )
        totals = self.db.execute(
            """
            SELECT COUNT(*), COALESCE(SUM(char_count),0),
                   COALESCE(SUM(page_count),0), COALESCE(SUM(ocr_pages),0),
                   COALESCE(SUM(embedded_pages),0)
            FROM document_features WHERE feature_version=?
            """,
            (FEATURE_VERSION,),
        ).fetchone()
        return {
            "feature_version": FEATURE_VERSION,
            "cached_documents": int(totals[0]),
            "text_characters": int(totals[1]),
            "pages": int(totals[2]),
            "ocr_pages": int(totals[3]),
            "embedded_text_pages": int(totals[4]),
            "status_counts": status_counts,
            "extractor_counts": extractor_counts,
        }

    def close(self) -> None:
        self.db.commit()
        self.db.close()


def load_documents(root: Path, examples_path: Path) -> dict[str, DocumentRef]:
    documents: dict[str, DocumentRef] = {}
    with examples_path.open(encoding="utf-8") as handle:
        for line in handle:
            if not line.strip():
                continue
            example = json.loads(line)
            for attachment in example.get("input", {}).get("attachments", []):
                download = attachment.get("download") or {}
                sha = download.get("sha256")
                rel = download.get("relative_path")
                if not sha or not rel or sha in documents:
                    continue
                documents[sha] = DocumentRef(
                    sha256=str(sha),
                    path=root / str(rel),
                    relative_path=str(rel),
                    mime=str(attachment.get("mimetype") or ""),
                    name=str(attachment.get("name") or Path(str(rel)).name),
                )
    return documents


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--root",
        default="/data/storage/odoo-training-dataset",
        help="Exported dataset root",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=0,
        help="Process at most N uncached documents (0 = all)",
    )
    parser.add_argument("--force", action="store_true")
    parser.add_argument("--progress-every", type=int, default=100)
    parser.add_argument("--ocr-timeout", type=int, default=90)
    return parser


def main() -> int:
    args = build_parser().parse_args()
    root = Path(args.root)
    examples_path = root / "training" / "examples.jsonl"
    features_dir = root / "features"
    db_path = features_dir / "document_features.sqlite3"
    manifest_path = features_dir / "document_features_manifest.json"

    if not examples_path.is_file():
        raise FileNotFoundError(examples_path)

    documents = load_documents(root, examples_path)
    print("FEATURE_VERSION =", FEATURE_VERSION)
    print("UNIQUE_DOCUMENTS =", len(documents))

    ocr = AdaptiveOcrEngine(timeout_seconds=args.ocr_timeout)
    registry = ExtractorRegistry(ocr)
    cache = FeatureCacheRepository(db_path)

    processed = 0
    reused = 0
    successes = 0
    unsupported = 0
    errors = 0
    missing = 0
    error_samples: list[tuple[str, str, str]] = []

    try:
        for document in documents.values():
            if not args.force and cache.is_current_success(document.sha256):
                reused += 1
                continue

            if args.limit and processed >= args.limit:
                break

            if not document.path.is_file():
                result = ExtractionResult(
                    status="error",
                    extractor="missing_file",
                    error="Source file does not exist",
                )
                cache.save(document, result)
                processed += 1
                errors += 1
                missing += 1
                continue

            extractor = registry.resolve(document.path)
            try:
                result = extractor.extract(document)
            except Exception as exc:
                result = ExtractionResult(
                    status="error",
                    extractor=extractor.__class__.__name__,
                    error=f"{type(exc).__name__}: {str(exc)[:1000]}",
                )

            cache.save(document, result)
            processed += 1

            if result.status == "success":
                successes += 1
            elif result.status == "unsupported":
                unsupported += 1
            else:
                errors += 1
                if len(error_samples) < 20:
                    error_samples.append(
                        (document.relative_path, result.extractor, result.error or "")
                    )

            if processed % 25 == 0:
                cache.commit()

            if args.progress_every and processed % args.progress_every == 0:
                print(
                    "PROGRESS",
                    f"processed={processed}",
                    f"success={successes}",
                    f"unsupported={unsupported}",
                    f"errors={errors}",
                    f"reused={reused}",
                    flush=True,
                )

        cache.commit()
        summary = cache.summary()
    finally:
        cache.close()

    features_dir.mkdir(parents=True, exist_ok=True)
    free_bytes = shutil.disk_usage(root).free
    db_size = db_path.stat().st_size if db_path.exists() else 0

    manifest = {
        "generated_at": utc_now(),
        "dataset_root": str(root),
        "unique_documents": len(documents),
        "this_run": {
            "processed": processed,
            "reused": reused,
            "successes": successes,
            "unsupported": unsupported,
            "errors": errors,
            "missing": missing,
        },
        "cache": summary,
        "feature_db_bytes": db_size,
        "free_bytes": free_bytes,
        "error_samples": error_samples,
    }
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    print("FEATURE_CACHE_DB =", db_path)
    print("FEATURE_CACHE_DB_MIB =", round(db_size / 1024**2, 2))
    print("FREE_GIB =", round(free_bytes / 1024**3, 3))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
